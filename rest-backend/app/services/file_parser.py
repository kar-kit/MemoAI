import io
from fastapi import UploadFile
from pypdf import PdfReader
from pptx import Presentation

MAX_CHARS = 12000


def _trim(text: str) -> str:
    text = (text or "").strip()
    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS] + "\n\n[...truncated]"
    return text


def extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    out = []
    for page in reader.pages:
        t = page.extract_text() or ""
        if t.strip():
            out.append(t)
    return _trim("\n\n".join(out))


def extract_text_from_pptx(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    out = []

    for i, slide in enumerate(prs.slides, start=1):
        out.append(f"[Slide {i}]")

        for shape in slide.shapes:
            # Most text boxes / titles
            if getattr(shape, "has_text_frame", False) and shape.text_frame:  # type: ignore
                t = shape.text_frame.text.strip()  # type: ignore
                if t:
                    out.append(t)

            # Tables
            if getattr(shape, "has_table", False) and shape.table:  # type: ignore
                table_lines = []
                for row in shape.table.rows:  # type: ignore
                    cells = [cell.text.strip() for cell in row.cells]
                    table_lines.append(" | ".join(cells))
                if table_lines:
                    out.append("\n".join(table_lines))

    return _trim("\n\n".join(out))


async def extract_text_from_file(file: UploadFile) -> str:
    filename = (file.filename or "").lower()
    data = await file.read()

    if filename.endswith(".pdf"):
        return extract_text_from_pdf(data)

    if filename.endswith(".pptx"):
        return extract_text_from_pptx(data)

    # IMPORTANT: .ppt (old binary) is NOT supported by python-pptx
    # Recommend users re-save as .pptx
    if filename.endswith(".ppt"):
        raise ValueError(
            "Old .ppt files aren’t supported. Please upload .pptx (re-save in PowerPoint)."
        )

    raise ValueError("Unsupported file type. Upload .pdf or .pptx.")
